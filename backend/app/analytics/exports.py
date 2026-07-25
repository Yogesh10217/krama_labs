import io
from typing import Dict, Any
from datetime import datetime

import openpyxl
from fpdf import FPDF

def generate_excel_report(data: Dict[str, Any], report_type: str) -> bytes:
    """Generate an Excel report from analytics data."""
    wb = openpyxl.Workbook()
    
    # Remove default sheet
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])
        
    for section_name, records in data.items():
        ws = wb.create_sheet(title=section_name[:31]) # Excel sheet name limit
        if not records:
            ws.append(["No data available for this section."])
            continue
            
        # Get headers from first record
        # Exclude SQLAlchemy internal state
        headers = [k for k in records[0].keys() if not k.startswith('_')]
        ws.append(headers)
        
        for record in records:
            row = []
            for h in headers:
                val = record.get(h)
                # Format datetime/date to string if needed, openpyxl usually handles datetime
                row.append(val)
            ws.append(row)
            
    # Save to memory
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()

def generate_pdf_report(data: Dict[str, Any], report_type: str) -> bytes:
    """Generate a PDF report from analytics data."""
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    
    # Title
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(0, 10, f"Analytics Report: {report_type.capitalize()}", ln=True, align='C')
    pdf.set_font("Arial", '', 10)
    pdf.cell(0, 10, f"Generated on: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')}", ln=True, align='C')
    pdf.ln(10)
    
    for section_name, records in data.items():
        pdf.set_font("Arial", 'B', 14)
        pdf.cell(0, 10, section_name.replace('_', ' ').capitalize(), ln=True)
        pdf.ln(5)
        
        if not records:
            pdf.set_font("Arial", 'I', 10)
            pdf.cell(0, 10, "No data available.", ln=True)
            pdf.ln(5)
            continue
            
        pdf.set_font("Arial", '', 10)
        headers = [k for k in records[0].keys() if not k.startswith('_') and k != 'id']
        
        # Simple rendering for PDF
        for i, record in enumerate(records):
            pdf.set_font("Arial", 'B', 10)
            pdf.cell(0, 8, f"Record #{i+1}", ln=True)
            pdf.set_font("Arial", '', 10)
            for h in headers:
                val = record.get(h)
                pdf.cell(0, 6, f"{h}: {val}", ln=True)
            pdf.ln(4)
            
        pdf.ln(10)
        
    return pdf.output(dest='S').encode('latin1')
