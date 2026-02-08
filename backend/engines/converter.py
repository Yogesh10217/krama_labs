"""Document converter - converts PDF, PPTX, Excel, Images to PIL Images."""

import os
import tempfile
from io import BytesIO
from pathlib import Path
from typing import List, Tuple

import fitz  # PyMuPDF
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
from PIL import Image
from pptx import Presentation

from config import Config
from models import DocumentType


class DocumentConverter:
    """Convert various document formats to images."""

    @staticmethod
    def detect_type(path: str) -> DocumentType:
        ext = Path(path).suffix.lower()
        mapping = {
            ".pdf": DocumentType.PDF,
            ".pptx": DocumentType.PPTX,
            ".ppt": DocumentType.PPTX,
            ".xlsx": DocumentType.EXCEL,
            ".xls": DocumentType.EXCEL,
            ".png": DocumentType.IMAGE,
            ".jpg": DocumentType.IMAGE,
            ".jpeg": DocumentType.IMAGE,
            ".bmp": DocumentType.IMAGE,
            ".tiff": DocumentType.IMAGE,
            ".webp": DocumentType.IMAGE,
            ".gif": DocumentType.IMAGE,
        }
        return mapping.get(ext, DocumentType.UNKNOWN)

    @staticmethod
    def ensure_rgb(img: Image.Image) -> Image.Image:
        if img.mode == "RGBA":
            background = Image.new("RGB", img.size, (255, 255, 255))
            background.paste(img, mask=img.split()[3])
            return background
        elif img.mode != "RGB":
            return img.convert("RGB")
        return img

    @classmethod
    def pdf_to_images(cls, path: str) -> List[Image.Image]:
        images = []
        doc = fitz.open(path)
        for page in doc:
            mat = fitz.Matrix(Config.DPI_FOR_PDF / 72, Config.DPI_FOR_PDF / 72)
            pix = page.get_pixmap(matrix=mat, alpha=False)
            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            if img.size[0] > Config.MAX_IMAGE_SIZE[0] or img.size[1] > Config.MAX_IMAGE_SIZE[1]:
                img.thumbnail(Config.MAX_IMAGE_SIZE, Image.Resampling.LANCZOS)
            images.append(img)
        doc.close()
        return images

    @classmethod
    def pptx_to_images(cls, path: str) -> List[Image.Image]:
        """Convert PPTX to images via LibreOffice (if available) or placeholder."""
        with tempfile.TemporaryDirectory() as tmpdir:
            try:
                import subprocess
                subprocess.run(
                    ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", tmpdir, path],
                    capture_output=True, timeout=120, check=True,
                )
                pdf_path = os.path.join(tmpdir, Path(path).stem + ".pdf")
                if os.path.exists(pdf_path):
                    return cls.pdf_to_images(pdf_path)
            except Exception:
                pass
        # Fallback: extract text from slides as images
        try:
            prs = Presentation(path)
            images = []
            for slide in prs.slides:
                fig, ax = plt.subplots(figsize=(16, 9))
                ax.set_xlim(0, 10)
                ax.set_ylim(0, 7.5)
                ax.axis("off")
                y = 7
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            text = para.text.strip()
                            if text:
                                ax.text(0.5, y, text, fontsize=10, va="top")
                                y -= 0.4
                buf = BytesIO()
                plt.savefig(buf, format="png", bbox_inches="tight", dpi=150, facecolor="white")
                buf.seek(0)
                images.append(Image.open(buf).copy())
                plt.close(fig)
                buf.close()
            return images if images else [Image.new("RGB", (1920, 1080), "white")]
        except Exception:
            return [Image.new("RGB", (1920, 1080), "white")]

    @classmethod
    def excel_to_images(cls, path: str) -> List[Image.Image]:
        images = []
        wb = openpyxl.load_workbook(path, data_only=True)
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows = [
                [str(c) if c is not None else "" for c in row]
                for row in sheet.iter_rows(values_only=True)
                if any(c is not None for c in row)
            ][:100]
            if not rows:
                continue
            max_cols = min(max(len(r) for r in rows), 20)
            fig, ax = plt.subplots(figsize=(max(12, max_cols * 1.2), max(6, len(rows) * 0.35)))
            ax.axis("off")
            ax.set_title(f"Sheet: {sheet_name}", fontsize=12, fontweight="bold")
            table_data = [r[:max_cols] for r in rows]
            table = ax.table(cellText=table_data, loc="center", cellLoc="left")
            table.auto_set_font_size(False)
            table.set_fontsize(8)
            table.scale(1.2, 1.5)
            buf = BytesIO()
            plt.savefig(buf, format="png", bbox_inches="tight", dpi=150, facecolor="white")
            buf.seek(0)
            images.append(Image.open(buf).copy())
            plt.close(fig)
            buf.close()
        wb.close()
        return images if images else [Image.new("RGB", (800, 600), "white")]

    @classmethod
    def load_image(cls, path: str) -> List[Image.Image]:
        img = Image.open(path)
        img = cls.ensure_rgb(img)
        if img.size[0] > Config.MAX_IMAGE_SIZE[0] or img.size[1] > Config.MAX_IMAGE_SIZE[1]:
            img.thumbnail(Config.MAX_IMAGE_SIZE, Image.Resampling.LANCZOS)
        return [img]

    @classmethod
    def convert(cls, file_path: str) -> Tuple[List[Image.Image], DocumentType]:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")
        doc_type = cls.detect_type(file_path)
        converters = {
            DocumentType.PDF: cls.pdf_to_images,
            DocumentType.PPTX: cls.pptx_to_images,
            DocumentType.EXCEL: cls.excel_to_images,
            DocumentType.IMAGE: cls.load_image,
        }
        converter = converters.get(doc_type)
        if not converter:
            raise ValueError(f"Unsupported document type: {doc_type}")
        images = converter(file_path)
        if not images:
            raise ValueError(f"No images extracted from: {file_path}")
        return images, doc_type
